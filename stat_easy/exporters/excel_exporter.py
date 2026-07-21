"""Excel(.xlsx) 出力。openpyxl で APA 風（三線表）スタイルを適用する。"""
from __future__ import annotations

from io import BytesIO

import pandas as pd


def export_tables(tables: dict[str, pd.DataFrame]) -> bytes:
    """シート名 -> DataFrame の辞書を 1 つの xlsx にまとめて出力。

    APA 風の三線表（上下に太線、ヘッダー下に細線）を適用する。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Border, Side, Alignment, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    thick = Side(style="thick", color="1A1A2E")
    thin = Side(style="thin", color="1A1A2E")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E6091")
    center = Alignment(horizontal="center", vertical="center")

    for sheet_name, df in tables.items():
        ws = wb.create_sheet(title=str(sheet_name)[:31])
        ws.append(list(df.columns))
        for _, row in df.iterrows():
            ws.append([_fmt(v) for v in row.tolist()])

        n_rows = len(df) + 1
        n_cols = len(df.columns)
        for c in range(1, n_cols + 1):
            hc = ws.cell(row=1, column=c)
            hc.font = header_font
            hc.fill = header_fill
            hc.alignment = center
            # 上罫線（太）+ ヘッダー下罫線（細）
            hc.border = Border(top=thick, bottom=thin)
            # 最終行の下罫線（太）
            bc = ws.cell(row=n_rows, column=c)
            bc.border = Border(bottom=thick)
        # 列幅自動調整
        for c in range(1, n_cols + 1):
            letter = ws.cell(row=1, column=c).column_letter
            maxlen = max([len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, n_rows + 1)])
            ws.column_dimensions[letter].width = min(40, max(10, maxlen + 2))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _fmt(v):
    if isinstance(v, float):
        return round(v, 4)
    return v
