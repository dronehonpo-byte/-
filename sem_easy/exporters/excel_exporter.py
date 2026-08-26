"""Excel(.xlsx) 出力。三線表風のスタイルを適用する。"""
from __future__ import annotations

from io import BytesIO

import pandas as pd


def export(tables: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    thick = Side(style="thick", color="1A1A2E")
    thin = Side(style="thin", color="1A1A2E")
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E4E79")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for name, df in tables.items():
        if df is None or len(df) == 0:
            continue
        ws = wb.create_sheet(title=str(name)[:31])
        ws.append([str(c) for c in df.columns])
        for _, row in df.iterrows():
            ws.append([_cell(v) for v in row.tolist()])
        ncols, nrows = len(df.columns), len(df) + 1
        for c in range(1, ncols + 1):
            hc = ws.cell(row=1, column=c)
            hc.font, hc.fill, hc.alignment = head_font, head_fill, center
            hc.border = Border(top=thick, bottom=thin)
            ws.cell(row=nrows, column=c).border = Border(bottom=thick)
            letter = hc.column_letter
            width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, nrows + 1))
            ws.column_dimensions[letter].width = min(60, max(12, width + 2))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, float):
        return round(v, 4)
    return v
