"""Word(.docx) 出力。図・表・解説・分析条件をまとめたレポート。"""
from __future__ import annotations

from io import BytesIO

import pandas as pd

METHODS_NOTE = (
    "本分析の推定は Python の SEM ライブラリ（semopy）および NumPy / SciPy により実行した。"
    "係数の計算に生成 AI は使用しておらず、同一のデータ・設定・乱数シードに対して"
    "常に同一の結果が得られる（再現性を担保）。"
    "なお、適合度が良好であることは、当該モデルが唯一正しいこと、"
    "または変数間の因果関係を証明することを意味しない。"
)


def export(result, tables: dict, images: list | None = None,
           heading: str = "SEM 分析レポート") -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.add_heading(heading, level=0)
    doc.add_paragraph(f"モデル：{result.model_name}")
    doc.add_paragraph(f"推定法：{result.estimator}　／　分析に用いた n = {result.n_used:,}")

    if images:
        doc.add_heading("パス図", level=1)
        for caption, png in images:
            doc.add_picture(BytesIO(png), width=Inches(6.0))
            cap = doc.add_paragraph(caption)
            cap.alignment = 1

    for name, df in tables.items():
        if df is None or len(df) == 0:
            continue
        doc.add_heading(str(name), level=1)
        t = doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        for i, c in enumerate(df.columns):
            cell = t.rows[0].cells[i]
            cell.text = str(c)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(9)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(row.tolist()):
                cells[i].text = _fmt(v)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

    doc.add_heading("分析方法に関する注記", level=1)
    doc.add_paragraph(METHODS_NOTE)

    doc.add_paragraph("")
    f = doc.add_paragraph("開発：株式会社Miyabee")
    f.alignment = 1

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _fmt(v) -> str:
    if isinstance(v, float):
        if v != v:
            return "—"
        return f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v)
